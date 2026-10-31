# =============================================================================
# Acacia EVs and NVs: spectrum on CaF2 with long acquisition (100x, 100 % laser
# power, 4 x 60 s) vs SERS spectrum on Ag/SiNWs (50x LWD, 10 % laser power,
# 4 x 1 s). Raw spectra, same style as the paper figures.
#
# INPUT (DATA_DIR):
#   EV on CaF2     : s1-EV_Acacia-dropCaF2_532nm_1800gr_100%_100x_BC50_60s_4a(.txt)
#   EV on Ag/SiNWs : S2-EV_ACACIA-AgSiNW_532nm_1800gr_BC200_50XLF_1s_4a_10%(.txt)
#   NV on both     : Book1.xlsx (columns X1, Y1 = NV on CaF2; X2, Y2 = NV on Ag/SiNWs)
# HOW TO RUN (PyCharm): pip install numpy matplotlib openpyxl, check DATA_DIR, Run.
# OUTPUT (in DATA_DIR): CaF2_LongAcq_vs_AgSiNW.png / .tiff (600 dpi)
#
# The four spectra differ ~20-fold in counts, so each is drawn with the same height
# and its own counts scale bar (right), so that the real intensities stay visible.
# =============================================================================
import sys
from pathlib import Path

import numpy as np
import matplotlib
import matplotlib.pyplot as plt
import openpyxl

DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r'C:\Users\Hira Aman\Downloads\comparison ev nv')
if len(sys.argv) > 1:
    matplotlib.use('Agg')

COL = {'EV': '#0B5A7E', 'NV': '#B5541C'}
CAF2_COL = {'EV': '#85ACBE', 'NV': '#DAA98D'}      # light shade of the sample colour
NAME = {'EV': 'EV Acacia', 'NV': 'Nanovesicles Acacia'}
COND = {'Ag': '50\u00d7, 10% laser power, 4 \u00d7 1 s', 'CaF2': '100\u00d7, 100% laser power, 4 \u00d7 60 s'}
BAR = {'Ag': 50, 'CaF2': 500}                       # counts shown by the scale bars
LOW, HIGH = (400, 1800), (2700, 3100)
HEIGHT, GAP = 1.0, 0.35                             # drawn height of each spectrum, gap between them
FIGSIZE = (8.5, 8.0)
FS_NAME, FS_AXIS, FS_TICK, FS_PANEL = 10, 13, 10, 13   # same as Fig 2a/2b

plt.rcParams.update({'font.family': 'sans-serif',
                     'font.sans-serif': ['Arial', 'Liberation Sans', 'DejaVu Sans'],
                     'font.size': FS_TICK})


# ----------------------------------------------------------------------------
# 1. FILES
# ----------------------------------------------------------------------------
def simple(name):
    return name.lower().replace(' ', '_').replace('%', '')


def find_txt(test, what):
    hits = [f for f in DATA_DIR.iterdir() if f.is_file() and f.suffix.lower() in ('.txt', '')
            and test(simple(f.name))]
    if len(hits) != 1:
        sys.exit(f'Expected one {what} file in {DATA_DIR}, found {len(hits)}: {[h.name for h in hits]}')
    print(f'{what}: {hits[0].name}')
    return np.loadtxt(hits[0])


spec = {('EV', 'CaF2'): find_txt(lambda n: 'ev_acacia-dropcaf2' in n and '100x' in n and '60s' in n,
                                 'EV on CaF2 (100x, 60 s)'),
        ('EV', 'Ag'): find_txt(lambda n: n.startswith('s2-ev_acacia-agsinw'), 'EV on Ag/SiNWs')}

books = [f for f in DATA_DIR.glob('*.xlsx') if not f.name.startswith('~$')]
for book in books:                                   # the workbook with the NV spectra
    ws = openpyxl.load_workbook(book, data_only=True).active
    head = [str(c or '') for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    if len(head) >= 4 and 'caf2' in head[1].lower() and 'agsinw' in head[3].lower():
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        spec[('NV', 'CaF2')] = np.array([(r[0], r[1]) for r in rows if r[0] is not None], float)
        spec[('NV', 'Ag')] = np.array([(r[2], r[3]) for r in rows if r[2] is not None], float)
        print(f'NV on CaF2 and on Ag/SiNWs: {book.name}\n  {head[1]}\n  {head[3]}')
        break
else:
    sys.exit(f'No workbook in {DATA_DIR} with NV on CaF2 (column B) and on Ag/SiNWs (column D)')


def view(d):
    x = d[:, 0]
    return ((x >= LOW[0]) & (x <= LOW[1])) | ((x >= HIGH[0]) & (x <= HIGH[1]))


# ----------------------------------------------------------------------------
# 2. FIGURE: (a) EVs, (b) NVs; in each, CaF2 below and Ag/SiNWs above
# ----------------------------------------------------------------------------
fig = plt.figure(figsize=FIGSIZE)
gs = fig.add_gridspec(2, 2, width_ratios=[LOW[1] - LOW[0], HIGH[1] - HIGH[0]], wspace=0.03, hspace=0.10)
for row, s in enumerate(['EV', 'NV']):
    axL = fig.add_subplot(gs[row, 0])
    axR = fig.add_subplot(gs[row, 1], sharey=axL)
    scale, offset = {}, {}
    for i, sub in enumerate(['CaF2', 'Ag']):
        d = spec[(s, sub)]
        y = d[view(d), 1]
        scale[sub] = HEIGHT / (y.max() - y.min())               # counts -> drawn units
        offset[sub] = i * (HEIGHT + GAP) - y.min() * scale[sub]
    for ax, (lo, hi) in [(axL, LOW), (axR, HIGH)]:
        for sub in ['CaF2', 'Ag']:
            d = spec[(s, sub)]
            m = (d[:, 0] >= lo) & (d[:, 0] <= hi)
            ax.plot(d[m, 0], d[m, 1] * scale[sub] + offset[sub], lw=0.6,
                    color=CAF2_COL[s] if sub == 'CaF2' else COL[s])
        ax.set_xlim(lo, hi)
        ax.tick_params(direction='in', length=5, labelsize=FS_TICK)
    axL.set_ylim(-0.25, 2 * HEIGHT + GAP + 0.30)
    # names and acquisition conditions, top left of each spectrum
    for i, sub in enumerate(['CaF2', 'Ag']):
        y_top = i * (HEIGHT + GAP) + HEIGHT
        where = 'CaF$_2$' if sub == 'CaF2' else 'Ag/SiNWs'
        axL.text(415, y_top + 0.03, f'{NAME[s]} on {where}  ({COND[sub]})', color=COL[s],
                 fontsize=FS_NAME, ha='left', va='bottom')
    # counts scale bar of each spectrum, in the flat part of the CH window (2700-2800 cm-1)
    for sub in ['CaF2', 'Ag']:
        d = spec[(s, sub)]
        k = (d[:, 0] >= 2700) & (d[:, 0] <= 2800)
        h = BAR[sub] * scale[sub]
        y0 = (d[k, 1] * scale[sub] + offset[sub]).max() + 0.06
        xb = 2715
        axR.plot([xb, xb], [y0, y0 + h], color='k', lw=1.5)
        axR.text(xb + 10, y0 + h / 2, f'{BAR[sub]}\ncounts', ha='left', va='center', fontsize=FS_NAME - 1)
    axL.spines['right'].set_visible(False)
    axR.spines['left'].set_visible(False)
    axL.tick_params(labelleft=False)
    axR.tick_params(left=False, right=True, labelleft=False)
    axL.set_xticks([400, 600, 800, 1000, 1200, 1400, 1600])
    axR.set_xticks([2800, 3000])
    if row == 0:
        axL.tick_params(labelbottom=False)
        axR.tick_params(labelbottom=False)
    dd = 0.012
    for ax, x0 in [(axL, 1), (axR, 0)]:                     # axis-break marks
        dx = dd * (HIGH[1] - HIGH[0]) / (LOW[1] - LOW[0]) if ax is axL else dd
        for y0 in (0, 1):
            ax.plot((x0 - dx, x0 + dx), (y0 - dd, y0 + dd), transform=ax.transAxes, color='k',
                    clip_on=False, lw=1.0)
    axL.set_ylabel('Intensity (a.u.)', fontsize=FS_AXIS)
    axL.text(-0.075, 1.0, f'({"ab"[row]})', transform=axL.transAxes, fontsize=FS_PANEL,
             ha='right', va='top')
fig.text(0.53, 0.012, 'Raman Shift (cm$^{-1}$)', ha='center', fontsize=FS_AXIS)
fig.subplots_adjust(left=0.08, right=0.98, bottom=0.075, top=0.98)


def save_figure(fig, path, **kw):
    """Save; if Windows keeps the old file open, save under a new name."""
    try:
        fig.savefig(path, **kw)
    except OSError:
        alt = path.with_name(path.stem + '_1' + path.suffix)
        fig.savefig(alt, **kw)
        print(f'{path.name} is open in another program; saved as {alt.name}')


save_figure(fig, DATA_DIR / 'CaF2_LongAcq_vs_AgSiNW.png', dpi=600)
save_figure(fig, DATA_DIR / 'CaF2_LongAcq_vs_AgSiNW.tiff', dpi=600, pil_kwargs={'compression': 'tiff_lzw'})
print(f'\nSaved in: {DATA_DIR}')
if len(sys.argv) <= 1:
    plt.show()
