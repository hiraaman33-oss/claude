"""Generate PLSDA_SVM_5fold_V3.docx from PLSDA_SVM_5fold_V3.py"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

SRC = "/home/user/claude/PLSDA_SVM_5fold_V3.py"
DST = "/home/user/claude/PLSDA_SVM_5fold_V3.docx"

doc = Document()

# ── title block ─────────────────────────────────────────────────────────────
title = doc.add_heading("PLS-DA + SVM-RBF  |  4-Class  |  Book1.xls  |  VIP > 1  |  5-fold CV", 0)
title.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

sub = doc.add_paragraph(
    "Raman Spectroscopy — 4-Class SSY Classification\n"
    "All 80 spectra (no train/test split)  |  5-fold Stratified CV only\n"
    "PLS-DA (5 LVs) + SVM-RBF (C=10, gamma=scale)  |  VIP selection: ALL variables with VIP > 1\n"
    "File: Book1.xls  |  Sheet: Sheet1  |  450-1800 cm⁻¹\n"
    "Path: C:/Users/Hira Aman/Desktop/PROF_DOMENICO'S"
)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in sub.runs:
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

doc.add_paragraph()

# ── how-to-run box ───────────────────────────────────────────────────────────
doc.add_heading("How to Run", 2)
htr = doc.add_paragraph()
htr.add_run("1. Install dependencies (one-time):\n").bold = True
htr.add_run("   pip install pandas openpyxl scikit-learn scipy matplotlib xlrd\n\n")
htr.add_run("2. Save the script as: PLSDA_SVM_5fold_V3.py\n").bold = True
htr.add_run(
    "   Place it anywhere convenient (e.g. Desktop).\n\n"
)
htr.add_run("3. Run:\n").bold = True
htr.add_run("   python PLSDA_SVM_5fold_V3.py\n\n")
htr.add_run("Outputs (1 folder):\n").bold = True
htr.add_run(
    "   PLSDA_SVM_VIPgt1\\\n"
    "   15 PNG figures + 1 Excel workbook (Results_5fold.xlsx)"
)
for run in htr.runs:
    run.font.size = Pt(10)

doc.add_paragraph()

# ── output file table ────────────────────────────────────────────────────────
doc.add_heading("Output Files", 2)
table = doc.add_table(rows=1, cols=2)
table.style = "Light List Accent 1"
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Filename"
hdr_cells[1].text = "Description"
for cell in hdr_cells:
    for run in cell.paragraphs[0].runs:
        run.bold = True

rows = [
    ("01_PLSDA_Scores_LV1_LV2.png",        "PLS-DA score plot LV1 x LV2 (full model, 5-fold CV misclassifications)"),
    ("02_PLSDA_Scores_LV1_LV3.png",        "PLS-DA score plot LV1 x LV3"),
    ("03_PLSDA_Scores_LV2_LV3.png",        "PLS-DA score plot LV2 x LV3"),
    ("04_PLSDA_Scores_3D.png",             "PLS-DA 3D score plot (LV1/LV2/LV3)"),
    ("05_PLSDA_Accuracy_Matrix_5fold.png", "PLS-DA accuracy confusion matrix — 5-fold CV (raw counts)"),
    ("06_PLSDA_Recall_Matrix_5fold.png",   "PLS-DA recall matrix — 5-fold CV (row-normalised, diagonal = recall)"),
    ("07_PLSDA_VIP_Scores.png",            "VIP scores with VIP > 1 threshold and selected variables shaded"),
    ("08_PLSDA_Loadings.png",              "PLS-DA loadings LV1/LV2/LV3 (VIP-selected variables)"),
    ("09_PLSDA_Permutation.png",           "Permutation test (999 permutations)"),
    ("10_SVM_Cluster_5fold.png",           "SVM-RBF cluster plot — 5-fold CV (PCA space)"),
    ("11_SVM_Accuracy_Matrix_5fold.png",   "SVM-RBF accuracy confusion matrix — 5-fold CV (raw counts)"),
    ("12_SVM_Recall_Matrix_5fold.png",     "SVM-RBF recall matrix — 5-fold CV (row-normalised, diagonal = recall)"),
    ("13_Accuracy_Comparison.png",         "Overall accuracy bar chart: PLS-DA vs SVM-RBF"),
    ("14_Per_Class_Recall_5fold.png",      "Per-class recall bar chart (5-fold CV, both models)"),
    ("15_Mean_Spectra.png",               "Mean Raman spectra ± 1 SD (all 80 spectra)"),
    ("Results_5fold.xlsx",                "Excel workbook — Summary, CMs, VIP scores, Loadings, Permutation"),
]
for fn, desc in rows:
    row = table.add_row().cells
    row[0].text = fn
    row[1].text = desc
    for cell in row:
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(9)

doc.add_paragraph()

# ── Python source code ───────────────────────────────────────────────────────
doc.add_heading("Python Script: PLSDA_SVM_5fold_V3.py", 1)

with open(SRC, encoding="utf-8") as f:
    code = f.read()

code_para = doc.add_paragraph()
code_para.paragraph_format.space_before = Pt(0)
code_para.paragraph_format.space_after  = Pt(0)

for line in code.split("\n"):
    run = code_para.add_run(line + "\n")
    run.font.name = "Courier New"
    run.font.size = Pt(8)

# ── save ─────────────────────────────────────────────────────────────────────
doc.save(DST)
print(f"Saved: {DST}")
